from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "Prototype Submission Template _ CoCo CLI Hackathon GCC Edition.pptx"
ASSETS = ROOT / "submission-assets"
OUTPUT = ASSETS / "SupplyGraph_AI_Submission_Deck.pptx"

WHITE = RGBColor(255, 255, 255)
MUTED = RGBColor(181, 201, 211)
TEAL = RGBColor(69, 188, 228)
GREEN = RGBColor(93, 195, 173)
INK = RGBColor(15, 24, 29)
PANEL = RGBColor(17, 29, 35)
LINE = RGBColor(52, 80, 92)


def remove_shape(shape: object) -> None:
    element = shape._element  # type: ignore[attr-defined]
    element.getparent().remove(element)


def add_text(
    slide: object,
    text: str,
    left: float,
    top: float,
    width: float,
    height: float,
    *,
    size: int = 16,
    color: RGBColor = WHITE,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
) -> object:
    box = slide.shapes.add_textbox(  # type: ignore[attr-defined]
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    return box


def add_panel(slide: object, left: float, top: float, width: float, height: float) -> object:
    panel = slide.shapes.add_shape(  # type: ignore[attr-defined]
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = PANEL
    panel.line.color.rgb = LINE
    panel.line.width = Pt(1)
    return panel


def add_bullets(
    slide: object,
    items: list[str],
    left: float,
    top: float,
    width: float,
    height: float,
    *,
    size: int = 13,
    color: RGBColor = WHITE,
) -> None:
    box = slide.shapes.add_textbox(  # type: ignore[attr-defined]
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = color
        paragraph.space_after = Pt(8)
        paragraph.text = f"•  {paragraph.text}"


def add_node(slide: object, label: str, left: float, top: float, width: float) -> object:
    node = slide.shapes.add_shape(  # type: ignore[attr-defined]
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(0.82),
    )
    node.fill.solid()
    node.fill.fore_color.rgb = PANEL
    node.line.color.rgb = TEAL
    node.line.width = Pt(1.4)
    frame = node.text_frame
    frame.clear()
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.text = label
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(13)
    paragraph.font.bold = True
    paragraph.font.color.rgb = WHITE
    return node


def connect(slide: object, start_x: float, end_x: float, y: float) -> None:
    connector = slide.shapes.add_connector(  # type: ignore[attr-defined]
        MSO_CONNECTOR.STRAIGHT,
        Inches(start_x),
        Inches(y),
        Inches(end_x),
        Inches(y),
    )
    connector.line.color.rgb = TEAL
    connector.line.width = Pt(2)
    connector.line.end_arrowhead = True


def crop_dashboard() -> Path:
    source = ASSETS / "dashboard-live-desktop.png"
    target = ASSETS / "dashboard-overview-crop.png"
    with Image.open(source) as image:
        crop_height = min(image.height, int(image.width / 1.62))
        image.crop((0, 0, image.width, crop_height)).save(target)
    return target


def clear_slide(slide: object) -> None:
    for shape in list(slide.shapes):  # type: ignore[attr-defined]
        if not shape.is_placeholder:
            remove_shape(shape)


def add_slide_heading(
    slide: object,
    number: int,
    title: str,
    subtitle: str,
) -> None:
    masthead = slide.shapes.add_shape(  # type: ignore[attr-defined]
        MSO_SHAPE.RECTANGLE,
        Inches(0),
        Inches(0),
        Inches(10),
        Inches(0.42),
    )
    masthead.fill.solid()
    masthead.fill.fore_color.rgb = INK
    masthead.line.fill.background()
    add_text(
        slide,
        "NOVA-AgenticIQ  |  SupplyGraph AI",
        0.35,
        0.08,
        3.4,
        0.2,
        size=8,
        color=WHITE,
        bold=True,
    )
    add_text(
        slide,
        "CoCo CLI HACKATHON  •  GCC Edition",
        6.3,
        0.08,
        3.35,
        0.2,
        size=8,
        color=TEAL,
        bold=True,
        align=PP_ALIGN.RIGHT,
    )
    add_text(
        slide,
        f"{number}  |  {title}",
        0.45,
        0.62,
        9.1,
        0.5,
        size=24,
        color=TEAL,
        bold=True,
    )
    add_text(slide, subtitle, 0.45, 1.06, 9.1, 0.35, size=13, color=MUTED)


def add_stat_card(
    slide: object,
    value: str,
    label: str,
    detail: str,
    left: float,
    top: float,
    width: float,
    *,
    accent: RGBColor = TEAL,
) -> None:
    add_panel(slide, left, top, width, 1.35)
    add_text(
        slide,
        value,
        left + 0.16,
        top + 0.12,
        width - 0.32,
        0.42,
        size=23,
        color=accent,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        label.upper(),
        left + 0.16,
        top + 0.53,
        width - 0.32,
        0.25,
        size=9,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        detail,
        left + 0.18,
        top + 0.81,
        width - 0.36,
        0.34,
        size=8,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )


def build_deck() -> None:
    presentation = Presentation(TEMPLATE)
    overview_image = crop_dashboard()

    while len(presentation.slides) < 10:
        presentation.slides.add_slide(presentation.slide_layouts[0])

    cover = presentation.slides[0]
    replacements = {
        "Team Name :": "Team Name: NOVA-AgenticIQ",
        "Problem Statement :": (
            "Problem Statement: Supply Chain Ontology and Governed Conversational Analytics"
        ),
        "Team Leader Name :": "Team Leader Name: Nikhilraj SV",
        "Team Size :": "Team Size: 1",
    }
    for shape in cover.shapes:
        if not hasattr(shape, "text_frame"):
            continue
        original = shape.text.strip()
        if original not in replacements:
            continue
        shape.text_frame.clear()
        paragraph = shape.text_frame.paragraphs[0]
        paragraph.text = replacements[original]
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(13 if "Problem" in original else 15)
        paragraph.font.bold = True
        paragraph.font.color.rgb = INK

    problem = presentation.slides[1]
    clear_slide(problem)
    add_slide_heading(
        problem,
        1,
        "Problem Brief",
        "One governed answer across planning, procurement, and logistics",
    )
    add_panel(problem, 0.45, 1.52, 4.05, 3.5)
    add_text(problem, "THE PAIN", 0.7, 1.68, 3.5, 0.3, size=11, color=TEAL, bold=True)
    add_bullets(
        problem,
        [
            "Supply-chain data is fragmented across suppliers, plants, shipments, orders, inventory, and quality.",
            "Teams calculate the same KPI differently, creating slow and low-trust decisions.",
            "Generic text-to-SQL can return plausible answers without an auditable definition or source.",
        ],
        0.7,
        2.02,
        3.55,
        2.1,
        size=12,
    )
    add_text(problem, "TARGET USERS", 0.7, 4.33, 3.5, 0.3, size=11, color=TEAL, bold=True)
    add_text(problem, "Planning  •  Procurement  •  Logistics", 0.7, 4.62, 3.5, 0.25, size=11)
    problem.shapes.add_picture(
        str(overview_image), Inches(4.72), Inches(1.52), width=Inches(4.78)
    )

    ontology = presentation.slides[2]
    clear_slide(ontology)
    add_slide_heading(
        ontology,
        2,
        "Supply-chain Ontology",
        "A canonical model connects operational cause to customer impact",
    )
    ontology_rows = [
        [
            ("Supplier", 0.55),
            ("Part", 2.75),
            ("Inventory", 4.95),
            ("Plant", 7.15),
        ],
        [
            ("Supplier", 0.55),
            ("Shipment", 2.75),
            ("Order", 4.95),
            ("Customer", 7.15),
        ],
    ]
    for row_index, row in enumerate(ontology_rows):
        top = 1.65 + row_index * 1.2
        for label, left in row:
            add_node(ontology, label, left, top, 1.55)
        for index in range(3):
            connect(ontology, row[index][1] + 1.55, row[index + 1][1], top + 0.41)
    add_node(ontology, "Quality Event", 3.85, 4.05, 2.05)
    add_text(ontology, "Supplier + Part", 1.45, 4.23, 2.1, 0.25, size=10, color=INK, bold=True, align=PP_ALIGN.RIGHT)
    connect(ontology, 3.6, 3.85, 4.46)
    add_panel(ontology, 6.35, 3.78, 2.35, 1.25)
    add_text(
        ontology,
        "ONE MODEL\nTHREE PERSONAS\nONE METRIC DEFINITION",
        6.58,
        3.95,
        1.9,
        0.82,
        size=11,
        color=GREEN,
        bold=True,
        align=PP_ALIGN.CENTER,
    )

    dashboard = presentation.slides[3]
    clear_slide(dashboard)
    add_slide_heading(
        dashboard,
        3,
        "Live Control Tower",
        "A live operational surface for network risk and performance",
    )
    dashboard.shapes.add_picture(
        str(overview_image), Inches(0.45), Inches(1.5), width=Inches(6.55)
    )
    add_stat_card(dashboard, "77.17%", "On-time delivery", "Completed shipments", 7.25, 1.48, 2.25)
    add_stat_card(dashboard, "96.15%", "Fill rate", "10,000 orders", 7.25, 2.84, 2.25, accent=GREEN)
    add_stat_card(dashboard, "554", "At-risk locations", "Part-plant exposure", 7.25, 4.2, 2.25, accent=RGBColor(242, 151, 92))

    architecture = presentation.slides[4]
    clear_slide(architecture)
    add_slide_heading(
        architecture,
        4,
        "Governed AI: Snowflake Cortex Architecture",
        "AI interprets the question; governance controls the answer",
    )
    nodes = [
        ("Business\nUser", 0.3),
        ("React\nUI", 1.9),
        ("FastAPI\nContract", 3.5),
        ("Cortex\nIntent", 5.1),
        ("Approved\nSQL", 6.7),
        ("Snowflake\nMetrics", 8.3),
    ]
    for label, left in nodes:
        add_node(architecture, label, left, 1.78, 1.35)
    for index in range(len(nodes) - 1):
        connect(architecture, nodes[index][1] + 1.35, nodes[index + 1][1], 2.19)
    add_panel(architecture, 0.5, 3.05, 4.2, 1.45)
    add_text(architecture, "WHAT CORTEX MAY RETURN", 0.78, 3.25, 3.65, 0.25, size=10, color=TEAL, bold=True)
    add_text(architecture, '{"metric_name": "ON_TIME_DELIVERY_RATE",\n "period_mode": "LATEST_MONTH"}', 0.78, 3.57, 3.65, 0.58, size=11)
    add_panel(architecture, 4.95, 3.05, 4.55, 1.45)
    add_text(architecture, "WHAT CORTEX CAN NEVER DO", 5.23, 3.25, 4.0, 0.25, size=10, color=RGBColor(242, 151, 92), bold=True)
    add_text(architecture, "Generate SQL  •  Change data  •  Bypass metric definitions", 5.23, 3.62, 4.0, 0.42, size=11)
    add_text(
        architecture,
        "Invalid or unavailable AI output falls back to deterministic intent matching.",
        0.7, 4.77, 8.6, 0.3,
        size=10,
        color=INK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )

    evidence = presentation.slides[5]
    clear_slide(evidence)
    add_slide_heading(
        evidence,
        5,
        "Evidence, Not Just an Answer",
        "Every answer carries its definition, formula, source, period, and SQL",
    )
    evidence.shapes.add_picture(
        str(ASSETS / "governed-answer-evidence.png"),
        Inches(0.5), Inches(1.5), width=Inches(9.0)
    )

    personas = presentation.slides[6]
    clear_slide(personas)
    add_slide_heading(
        personas,
        6,
        "Persona Consistency",
        "Recommendations adapt by role; metric truth never changes",
    )
    persona_cards = [
        ("PLANNING", "Adjust supply coverage and priorities."),
        ("PROCUREMENT", "Focus supplier and cost discussions."),
        ("LOGISTICS", "Prioritize delivery-flow investigation."),
    ]
    for index, (label, recommendation) in enumerate(persona_cards):
        left = 0.45 + index * 3.05
        add_panel(personas, left, 1.62, 2.72, 2.05)
        add_text(personas, label, left + 0.2, 1.84, 2.32, 0.3, size=10, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
        add_text(personas, "56.91%", left + 0.2, 2.19, 2.32, 0.48, size=25, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
        add_text(personas, recommendation, left + 0.28, 2.82, 2.16, 0.5, size=10, color=WHITE, align=PP_ALIGN.CENTER)
    add_panel(personas, 0.75, 4.02, 8.5, 0.92)
    add_text(
        personas,
        "IDENTICAL ACROSS ALL PERSONAS: metric  •  definition  •  evidence  •  SQL  •  result rows",
        1.05, 4.27, 7.9, 0.34,
        size=11, color=GREEN, bold=True, align=PP_ALIGN.CENTER,
    )

    scale = presentation.slides[7]
    clear_slide(scale)
    add_slide_heading(
        scale,
        7,
        "Impact and Scale",
        "A realistic governed network with measurable operational signals",
    )
    scale_stats = [
        ("50", "Suppliers", "GCC network"),
        ("10", "Plants", "Distribution footprint"),
        ("10,000", "Orders", "Customer demand"),
        ("554", "Risk locations", "Part-plant exposure"),
        ("49", "Affected orders", "Delayed inbound supply"),
        ("$256.2M", "Landed cost", "Material + logistics"),
    ]
    for index, (value, label, detail) in enumerate(scale_stats):
        row = index // 3
        column = index % 3
        add_stat_card(
            scale,
            value,
            label,
            detail,
            0.55 + column * 3.05,
            1.62 + row * 1.58,
            2.72,
            accent=GREEN if row else TEAL,
        )
    add_text(scale, "All values returned from live governed Snowflake objects and cross-checked through the API.", 0.7, 4.78, 8.6, 0.25, size=9, color=INK, bold=True, align=PP_ALIGN.CENTER)

    engineering = presentation.slides[8]
    clear_slide(engineering)
    add_slide_heading(
        engineering,
        8,
        "Engineering Readiness",
        "Trust is enforced in code, Snowflake, tests, and deployment",
    )
    readiness = [
        ("12", "backend tests", "Guardrails + contracts"),
        ("1", "live browser flow", "API-to-UI proof"),
        ("3", "personas", "Canonical consistency"),
        ("0", "generated SQL", "Allowlisted queries only"),
    ]
    for index, (value, label, detail) in enumerate(readiness):
        add_stat_card(engineering, value, label, detail, 0.45 + index * 2.32, 1.55, 2.08)
    add_panel(engineering, 0.45, 3.2, 4.35, 1.55)
    add_text(engineering, "SNOWFLAKE SECURITY", 0.75, 3.42, 3.75, 0.28, size=10, color=TEAL, bold=True)
    add_bullets(engineering, ["Read-only SUPPLYGRAPH_APP role", "Query tag: SUPPLYGRAPH_AI", "Credentials only in runtime secrets"], 0.75, 3.72, 3.75, 0.78, size=9)
    add_panel(engineering, 5.05, 3.2, 4.5, 1.55)
    add_text(engineering, "DEPLOYMENT", 5.35, 3.42, 3.9, 0.28, size=10, color=TEAL, bold=True)
    add_bullets(engineering, ["Multi-stage Docker image", "React + FastAPI on one origin", "Render blueprint + health check"], 5.35, 3.72, 3.9, 0.78, size=9)

    close = presentation.slides[9]
    clear_slide(close)
    add_slide_heading(
        close,
        9,
        "Thank You",
        "Why SupplyGraph AI wins: one network, one metric truth, every answer evidenced",
    )
    close_cards = [
        ("GOVERNED", "Definitions, formulas, source objects, SQL, and rows travel with the answer."),
        ("OPERATIONAL", "Supplier, part, plant, shipment, order, and customer impact connect in one workflow."),
        ("EXTENSIBLE", "New metrics plug into the catalog without weakening the trust boundary."),
    ]
    for index, (label, detail) in enumerate(close_cards):
        left = 0.45 + index * 3.05
        add_panel(close, left, 1.65, 2.72, 2.25)
        add_text(close, label, left + 0.22, 1.95, 2.28, 0.32, size=12, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
        add_text(close, detail, left + 0.3, 2.45, 2.12, 1.05, size=11, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(close, "NOVA-AgenticIQ  |  Nikhilraj SV  |  CoCo CLI Hackathon 2026", 0.75, 4.35, 8.5, 0.34, size=11, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_text(close, "One network. One metric truth. Every answer evidenced.", 0.75, 4.72, 8.5, 0.34, size=13, color=TEAL, bold=True, align=PP_ALIGN.CENTER)

    presentation.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    build_deck()